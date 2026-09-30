"""把快照渲染成 Excel 报表。"""

from __future__ import annotations

import math
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional, Sequence

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
BODY_FONT = Font(size=10, bold=True)
LINK_FONT = Font(size=10, bold=True, color="0563C1", underline="single")
_SIDE = Side(style="thin", color="D9D9D9")
BORDER = Border(left=_SIDE, right=_SIDE, top=_SIDE, bottom=_SIDE)

COLOR_MUTUAL = "2E7D32"    # 互关 - 绿
COLOR_FANS = "1565C0"      # 粉丝 - 蓝
COLOR_FOLLOWING = "6A1B9A"  # 关注 - 紫
COLOR_LOST = "C62828"      # 取关了我 - 红
COLOR_NEW = "00897B"       # 新粉丝 - 青
COLOR_LEFT = "EF6C00"      # 我取关了 - 橙
COLOR_HISTORY = "455A64"   # 概览 - 灰

def fmt_date(ts: Any) -> str:
    """秒级时间戳 -> ``YYYY-MM-DD``；无效值返回空串。"""
    try:
        n = int(ts or 0)
    except (TypeError, ValueError):
        return ""
    if n <= 0:
        return ""
    try:
        return datetime.fromtimestamp(n).strftime("%Y-%m-%d")
    except (ValueError, OSError, OverflowError):
        return ""


# 用户信息表的列：(表头, 取值函数)
USER_COLUMNS = [
    ("备注", lambda u: u.get("remark", "")),
    ("昵称", lambda u: u.get("nickname", "")),
    ("抖音号", lambda u: u.get("douyin_id", "")),
    ("关系", lambda u: u.get("relation", "")),
    ("账号注册时间", lambda u: fmt_date(u.get("created_time"))),
    ("粉丝数", lambda u: u.get("follower_count", 0)),
    ("关注数", lambda u: u.get("following_count", 0)),
    ("作品数", lambda u: u.get("aweme_count", 0)),
    ("认证", lambda u: u.get("custom_verify", "")),
    ("简介", lambda u: u.get("signature", "")),
]

# 首列是稳定的行号：按粉丝数等排完序后，再按它升序排一次就能回到
# 原始的时间顺序，不用撤销操作。列里的位置在下面按表头名查，不写死数字。
SEQ_HEADER = "序号"

WRAP_MAX_WIDTH = 46        # 折行列的列宽上限（再宽就换不了行了）
ROW_HEIGHT_PER_LINE = 15   # 折行后每行文字占的行高（磅）


def _display_width(value: Any) -> int:
    """中文按 2 个字符宽计算。"""
    return sum(2 if ord(ch) > 127 else 1 for ch in str(value))


def _longest_line_width(value: Any) -> int:
    """多行文本里最长那一行的宽度（折行列按它定列宽，而不是按整段长度）。"""
    return max((_display_width(line) for line in str(value).split("\n")), default=0)


def _estimated_lines(value: Any, width: float) -> int:
    """估算一个单元格折行后占几行。

    Excel/openpyxl 都不会自己算这个，不显式设行高的话行只会显示一行高，
    多出来的内容被截断 —— 而 WPS / LibreOffice 也不会自动撑开。
    """
    usable = max(width - 1.5, 1)
    total = 0
    for line in str(value).split("\n"):
        total += max(1, math.ceil(_display_width(line) / usable))
    return max(total, 1)


def write_sheet(
    ws,
    headers: Sequence[str],
    rows: Sequence[Sequence[Any]],
    color: str,
    links: Optional[Sequence[Optional[str]]] = None,
    link_col: Optional[int] = None,
    wrap_cols: Sequence[int] = (),
    left_cols: Sequence[int] = (),
) -> None:
    """写一个表：表头着色、冻结首行、加筛选、按内容自适应列宽。

    ``wrap_cols`` 里的列（1 起）开启自动折行并保留原文的分行，行高按内容估算；
    ``left_cols`` 里的列靠左显示，其余一律居中。
    """
    wrap_cols = tuple(wrap_cols)
    left_cols = tuple(left_cols)
    ws.append(list(headers))

    for idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=idx)
        cell.font = HEADER_FONT
        cell.fill = PatternFill("solid", fgColor=color)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER
    ws.freeze_panes = "A2"

    for r_idx, row in enumerate(rows, start=2):
        ws.append(list(row))
        for c_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=r_idx, column=c_idx)
            cell.font = BODY_FONT
            cell.border = BORDER
            wrapped = c_idx in wrap_cols
            # 折行的列用顶部对齐，免得行高略有富余时上下都空一截
            cell.alignment = Alignment(
                horizontal="left" if c_idx in left_cols else "center",
                vertical="top" if wrapped else "center",
                wrap_text=wrapped,
            )
        if links and link_col:
            url = links[r_idx - 2] if r_idx - 2 < len(links) else None
            if url:
                cell = ws.cell(row=r_idx, column=link_col)
                cell.hyperlink = url
                cell.font = LINK_FONT

    # 列宽：表头与内容取最大值。折行列按「最长的一行」算，并收在可读宽度内，
    # 否则一段 166 字的简介会把列撑到极限、反而不换行。
    widths: List[float] = [8.0] * (len(headers) + 1)
    for c_idx, header in enumerate(headers, start=1):
        if header == SEQ_HEADER:
            width = 6.0  # 行号列固定窄一点，不占地方
        elif c_idx in wrap_cols:
            width = _display_width(header) + 4
            for row in rows[:400]:
                if c_idx - 1 < len(row):
                    width = max(width, _longest_line_width(row[c_idx - 1]) + 3)
            width = min(max(width, 20), WRAP_MAX_WIDTH)
        else:
            width = _display_width(header) + 4
            for row in rows[:400]:
                if c_idx - 1 < len(row):
                    width = max(width, _display_width(row[c_idx - 1]) + 3)
            width = min(max(width, 8), 52)
        widths[c_idx] = width
        ws.column_dimensions[get_column_letter(c_idx)].width = width

    # 折行列显式设行高 —— 不设的话行只显示一行高，多出来的文字会被截掉
    if wrap_cols:
        for r_idx, row in enumerate(rows, start=2):
            lines = 1
            for c_idx in wrap_cols:
                if c_idx - 1 < len(row):
                    lines = max(lines, _estimated_lines(row[c_idx - 1], widths[c_idx]))
            if lines > 1:
                ws.row_dimensions[r_idx].height = min(lines * ROW_HEIGHT_PER_LINE, 409)

    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"


def _col(headers: Sequence[str], name: str) -> int:
    """按表头名取列号（1 起）。改列顺序时不用手动改一堆魔法数字。"""
    return headers.index(name) + 1


def _user_rows(users: Sequence[Dict[str, Any]]):
    headers = [SEQ_HEADER] + [c[0] for c in USER_COLUMNS]
    rows = [
        [i] + [fn(u) for _, fn in USER_COLUMNS]
        for i, u in enumerate(users, start=1)
    ]
    links = [u.get("home_url") or None for u in users]
    return headers, rows, links


def _write_user_sheet(wb, name: str, color: str, users: Sequence[Dict[str, Any]]) -> None:
    """互关/粉丝/关注三张表结构一样，统一在这里写（简介列开启折行）。"""
    headers, rows, links = _user_rows(users)
    signature = _col(headers, "简介")
    write_sheet(
        wb.create_sheet(name), headers, rows, color,
        links, link_col=_col(headers, "抖音号"),
        wrap_cols=(signature,),
        # 简介整段靠左读起来才顺，其余列居中
        left_cols=(signature,),
    )


def _event_rows(events: Sequence[Dict[str, Any]], date_label: str):
    headers = [SEQ_HEADER, "备注", "昵称", "抖音号", date_label, "粉丝数"]
    rows = [
        [
            i,
            e.get("remark", ""),
            e.get("nickname", ""),
            e.get("douyin_id", ""),
            e.get("date", ""),
            e.get("follower_count", 0),
        ]
        for i, e in enumerate(events, start=1)
    ]
    return headers, rows


def build_workbook(
    *,
    mutual: List[Dict[str, Any]],
    fans: List[Dict[str, Any]],
    following: List[Dict[str, Any]],
    lost_fans: List[Dict[str, Any]],
    new_fans: List[Dict[str, Any]],
    i_unfollowed: List[Dict[str, Any]],
    runs_summary: List[Dict[str, Any]],
    has_fans: bool,
    has_following: bool,
    warnings: List[str],
    out_path,
) -> Path:
    """生成报表并保存，返回实际写入路径。"""
    wb = Workbook()
    wb.remove(wb.active)

    _write_user_sheet(wb, "互关好友", COLOR_MUTUAL, mutual)

    if has_fans:
        _write_user_sheet(wb, "我的粉丝", COLOR_FANS, fans)

    if has_following:
        _write_user_sheet(wb, "我的关注", COLOR_FOLLOWING, following)

    if lost_fans:
        headers, rows = _event_rows(lost_fans, "取关于")
        write_sheet(wb.create_sheet("取关了我"), headers, rows, COLOR_LOST)

    if new_fans:
        headers, rows = _event_rows(new_fans, "首次发现")
        write_sheet(wb.create_sheet("新粉丝"), headers, rows, COLOR_NEW)

    if i_unfollowed:
        headers, rows = _event_rows(i_unfollowed, "我取关于")
        write_sheet(wb.create_sheet("我取关了"), headers, rows, COLOR_LEFT)

    overview_headers = ["采集日期", "关注数", "粉丝数", "互关数", "新增粉丝", "取关了我", "我新关注", "我取关"]
    overview_rows = [
        [
            r.get("date", ""),
            r.get("counts", {}).get("following", 0),
            r.get("counts", {}).get("fans", 0),
            r.get("counts", {}).get("mutual", 0),
            r.get("delta", {}).get("new_fans", ""),
            r.get("delta", {}).get("unfollowed_me", ""),
            r.get("delta", {}).get("i_followed", ""),
            r.get("delta", {}).get("i_unfollowed", ""),
        ]
        for r in runs_summary
    ]
    write_sheet(wb.create_sheet("历史概览"), overview_headers, overview_rows, COLOR_HISTORY)

    if warnings:
        # 说明都是整句中文，和「简介」列一样折行 + 靠左，不要一行走到底
        ws = wb.create_sheet("提示")
        write_sheet(ws, ["说明"], [[w] for w in warnings], COLOR_HISTORY,
                    wrap_cols=(1,), left_cols=(1,))

    out = Path(out_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    suppress_text_number_warnings(out)
    return out


def suppress_text_number_warnings(path) -> None:
    """让 Excel 不再给「数字以文本形式存储」的格子画绿色小三角。

    抖音号有不少是纯数字（如 ``1564072609``），但它本质是 ID 不是数值，
    必须按文本存 —— 于是 Excel 会在每个这种格子上报「数字存储为文本」。
    这里往工作表 XML 里写 ``<ignoredErrors>``，等价于右键「忽略错误」。

    为什么动 XML：openpyxl 定义了 ``IgnoredErrors`` 类，却没挂到 Worksheet 上，
    没有 API 可用。按 schema，``ignoredErrors`` 必须排在 ``pageMargins`` 之后、
    ``drawing`` 之前，而 openpyxl 的输出里 ``pageMargins`` 正是最后一个元素，
    所以直接插在 ``</worksheet>`` 前面就是合法位置。
    """
    import re
    import zipfile

    path = Path(path)
    tmp = path.with_name(path.name + ".tmp")
    with zipfile.ZipFile(path) as src:
        with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as dst:
            for name in src.namelist():
                data = src.read(name)
                if re.fullmatch(r"xl/worksheets/sheet\d+\.xml", name):
                    xml = data.decode("utf-8")
                    if "<ignoredErrors" not in xml:
                        found = re.search(r'<dimension ref="([^"]+)"', xml)
                        ref = found.group(1) if found else "A1"
                        if ":" not in ref:
                            ref = f"A1:{ref}"
                        xml = xml.replace(
                            "</worksheet>",
                            '<ignoredErrors><ignoredError sqref="{}" '
                            'numberStoredAsText="1"/></ignoredErrors></worksheet>'.format(ref),
                        )
                        data = xml.encode("utf-8")
                dst.writestr(name, data)
    tmp.replace(path)


def harvest_remarks(xlsx_path) -> Dict[str, str]:
    """从已有报表里回收用户手填的「备注」，返回 ``{抖音号: 备注}``。

    同一个人会出现在「互关好友 / 我的粉丝 / 我的关注」多个表里，用户可能在任何
    一个表里填备注，所以按并集回收：**非空值优先**，空值只在该人从未出现过非空
    备注时占位（这样既不会让空单元格盖掉已填的备注，被清空的人也不会残留旧值）。

    要彻底删除某个人的备注，请删掉 ``remarks.csv`` 里对应行。
    """
    if not xlsx_path:
        return {}
    p = Path(xlsx_path)
    if not p.exists():
        return {}
    out: Dict[str, str] = {}
    try:
        wb = load_workbook(p, read_only=True, data_only=True)
    except Exception:
        return {}
    try:
        for name in ("互关好友", "我的粉丝", "我的关注"):
            if name not in wb.sheetnames:
                continue
            ws = wb[name]
            rows = ws.iter_rows(values_only=True)
            try:
                header = next(rows)
            except StopIteration:
                continue
            try:
                remark_i = header.index("备注")
                douyin_i = header.index("抖音号")
            except ValueError:
                continue
            for row in rows:
                if douyin_i >= len(row) or remark_i >= len(row):
                    continue
                did = row[douyin_i]
                if not did:
                    continue
                key = str(did).strip()
                raw_remark = row[remark_i]
                remark = str(raw_remark).strip() if raw_remark is not None else ""
                if remark:
                    out[key] = remark          # 非空值优先，覆盖先前的空占位
                else:
                    out.setdefault(key, "")    # 空值不覆盖已回收到的备注
    finally:
        wb.close()
    return out
